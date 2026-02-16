import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, numbers
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.models.financial_statement import FinancialStatement
from app.models.multiple import Multiple
from app.models.quote import Quote
from app.models.security import Security
from app.models.statement_item import StatementItem
from app.models.statement_item_dict import StatementItemDict


class ExportService:
    HEADER_FONT = Font(bold=True, size=11)
    HEADER_ALIGNMENT = Alignment(horizontal="center", wrap_text=True)
    NUMBER_FORMAT = '#,##0.00'

    def __init__(self, db: AsyncSession):
        self.db = db

    async def export_financials(
        self, company_id: int, statement_ids: list[int]
    ) -> BytesIO:
        wb = Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        for statement_id in statement_ids:
            stmt = (
                select(FinancialStatement)
                .where(
                    FinancialStatement.id == statement_id,
                    FinancialStatement.company_id == company_id,
                )
                .options(
                    selectinload(FinancialStatement.items).selectinload(
                        StatementItem.item_dict
                    )
                )
            )
            result = await self.db.execute(stmt)
            statement = result.scalar_one_or_none()
            if statement is None:
                continue

            sheet_name = (
                f"{statement.statement_type}_{statement.period_end.isoformat()}"
            )
            # Excel sheet names max 31 chars
            ws = wb.create_sheet(title=sheet_name[:31])

            # Headers
            headers = ["Code", "Name", "Value"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = self.HEADER_FONT
                cell.alignment = self.HEADER_ALIGNMENT

            # Metadata row
            ws.cell(row=2, column=1, value="Standard")
            ws.cell(row=2, column=2, value=statement.standard)
            ws.cell(row=3, column=1, value="Period")
            ws.cell(
                row=3,
                column=2,
                value=f"{statement.period_start} - {statement.period_end}",
            )
            ws.cell(row=4, column=1, value="Currency")
            ws.cell(row=4, column=2, value=statement.currency)
            ws.cell(row=5, column=1, value="Unit (x)")
            ws.cell(row=5, column=2, value=statement.unit_multiplier)

            # Items
            sorted_items = sorted(
                statement.items, key=lambda i: i.item_dict.sort_order
            )
            for row_idx, item in enumerate(sorted_items, 7):
                ws.cell(row=row_idx, column=1, value=item.item_dict.code)
                ws.cell(row=row_idx, column=2, value=item.item_dict.name_ru)
                value_cell = ws.cell(
                    row=row_idx,
                    column=3,
                    value=float(item.value) if item.value is not None else None,
                )
                if item.value is not None:
                    value_cell.number_format = self.NUMBER_FORMAT

            # Adjust column widths
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 50
            ws.column_dimensions["C"].width = 20

        # If no sheets were created, add a blank one
        if not wb.sheetnames:
            wb.create_sheet(title="No Data")

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    async def export_multiples(self, company_id: int) -> BytesIO:
        stmt = (
            select(Multiple)
            .where(Multiple.company_id == company_id)
            .order_by(Multiple.calc_date.desc())
        )
        result = await self.db.execute(stmt)
        multiples = result.scalars().all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Multiples"

        headers = [
            "Calc Date",
            "Period End",
            "P/E",
            "EV/EBITDA",
            "EV/Sales",
            "P/B",
            "P/S",
            "ROE",
            "ROA",
            "Debt/EBITDA",
            "Dividend Yield",
            "Market Cap",
            "Enterprise Value",
            "Currency",
        ]

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self.HEADER_FONT
            cell.alignment = self.HEADER_ALIGNMENT

        for row_idx, m in enumerate(multiples, 2):
            ws.cell(row=row_idx, column=1, value=m.calc_date.isoformat())
            ws.cell(row=row_idx, column=2, value=m.period_end.isoformat())
            for col_idx, attr in enumerate(
                [
                    "pe",
                    "ev_ebitda",
                    "ev_sales",
                    "pb",
                    "ps",
                    "roe",
                    "roa",
                    "debt_ebitda",
                    "dividend_yield",
                    "market_cap",
                    "enterprise_value",
                ],
                3,
            ):
                val = getattr(m, attr)
                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=float(val) if val is not None else None,
                )
                if val is not None:
                    cell.number_format = self.NUMBER_FORMAT
            ws.cell(row=row_idx, column=14, value=m.currency)

        # Adjust column widths
        for col_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    async def export_quotes(
        self,
        security_id: int,
        from_date: datetime.date | None = None,
        to_date: datetime.date | None = None,
    ) -> BytesIO:
        stmt = (
            select(Quote)
            .where(Quote.security_id == security_id)
            .order_by(Quote.trade_date)
        )
        if from_date:
            stmt = stmt.where(Quote.trade_date >= from_date)
        if to_date:
            stmt = stmt.where(Quote.trade_date <= to_date)

        result = await self.db.execute(stmt)
        quotes = result.scalars().all()

        # Get security info
        sec_result = await self.db.execute(
            select(Security).where(Security.id == security_id)
        )
        security = sec_result.scalar_one_or_none()
        ticker = security.ticker if security else str(security_id)

        wb = Workbook()
        ws = wb.active
        ws.title = f"Quotes {ticker}"[:31]

        headers = ["Date", "Open", "High", "Low", "Close", "Volume"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self.HEADER_FONT
            cell.alignment = self.HEADER_ALIGNMENT

        for row_idx, q in enumerate(quotes, 2):
            ws.cell(row=row_idx, column=1, value=q.trade_date.isoformat())
            for col_idx, attr in enumerate(["open", "high", "low", "close"], 2):
                val = getattr(q, attr)
                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=float(val) if val is not None else None,
                )
                if val is not None:
                    cell.number_format = self.NUMBER_FORMAT
            ws.cell(row=row_idx, column=6, value=q.volume)

        ws.column_dimensions["A"].width = 14
        for col_letter in ["B", "C", "D", "E"]:
            ws.column_dimensions[col_letter].width = 14
        ws.column_dimensions["F"].width = 16

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    async def generate_upload_template(self) -> BytesIO:
        wb = Workbook()
        wb.remove(wb.active)

        sheet_configs = [
            ("Balance Sheet", "balance_sheet"),
            ("Income Statement", "income_statement"),
            ("Cash Flow", "cash_flow"),
        ]

        for sheet_name, statement_type in sheet_configs:
            ws = wb.create_sheet(title=sheet_name)

            # Headers
            headers = ["Code", "Name (RU)", "Value"]
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = self.HEADER_FONT
                cell.alignment = self.HEADER_ALIGNMENT

            # Fetch dictionary items for this statement type
            stmt = (
                select(StatementItemDict)
                .where(StatementItemDict.statement_type == statement_type)
                .order_by(StatementItemDict.sort_order)
            )
            result = await self.db.execute(stmt)
            dict_items = result.scalars().all()

            for row_idx, item in enumerate(dict_items, 2):
                ws.cell(row=row_idx, column=1, value=item.code)
                ws.cell(row=row_idx, column=2, value=item.name_ru)
                # Value column left empty for user input

            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 55
            ws.column_dimensions["C"].width = 20

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
