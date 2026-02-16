import datetime
import io
import logging
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, numbers
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.financial_statement import FinancialStatement
from app.models.statement_item import StatementItem
from app.models.statement_item_dict import StatementItemDict

logger = logging.getLogger(__name__)


class ExcelTemplateHandler:
    """Handles generation and parsing of Excel upload templates."""

    def __init__(self, db: AsyncSession):
        self.db = db

    async def generate_template(self) -> io.BytesIO:
        """Generate a downloadable Excel template for financial data upload."""
        stmt = select(StatementItemDict).order_by(
            StatementItemDict.statement_type, StatementItemDict.sort_order
        )
        result = await self.db.execute(stmt)
        items = result.scalars().all()

        wb = Workbook()
        wb.remove(wb.active)

        sheet_map = {
            "balance_sheet": "Баланс",
            "income_statement": "Отчёт о прибылях и убытках",
            "cash_flow": "Отчёт о движении ДС",
        }

        header_font = Font(bold=True, size=11)
        header_align = Alignment(horizontal="center", wrap_text=True)

        for stmt_type, sheet_title in sheet_map.items():
            ws = wb.create_sheet(title=sheet_title)

            ws.cell(row=1, column=1, value="Код статьи").font = header_font
            ws.cell(row=1, column=2, value="Наименование").font = header_font
            ws.cell(row=1, column=3, value="Период 1").font = header_font
            ws.cell(row=1, column=4, value="Период 2").font = header_font
            ws.cell(row=1, column=5, value="Период 3").font = header_font

            for col in range(1, 6):
                ws.cell(row=1, column=col).alignment = header_align

            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 45
            ws.column_dimensions["C"].width = 18
            ws.column_dimensions["D"].width = 18
            ws.column_dimensions["E"].width = 18

            row_num = 2
            type_items = [i for i in items if i.statement_type == stmt_type]
            for item in type_items:
                ws.cell(row=row_num, column=1, value=item.code)
                ws.cell(row=row_num, column=2, value=item.name_ru)

                is_parent = any(i.parent_code == item.code for i in type_items)
                if is_parent:
                    ws.cell(row=row_num, column=1).font = Font(bold=True)
                    ws.cell(row=row_num, column=2).font = Font(bold=True)

                for col in range(3, 6):
                    ws.cell(row=row_num, column=col).number_format = '#,##0'

                row_num += 1

            # Add instructions row
            row_num += 1
            ws.cell(row=row_num, column=1, value="Инструкция:").font = Font(bold=True, color="0000FF")
            ws.cell(
                row=row_num + 1, column=1,
                value="Заполните значения в столбцах 'Период'. Единица измерения: тыс. руб."
            )
            ws.cell(
                row=row_num + 2, column=1,
                value="Укажите даты периодов в ячейках заголовков (например, 31.12.2024)."
            )

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    async def parse_upload(
        self,
        file_content: bytes,
        company_id: int,
        standard: str = "RSBU",
    ) -> dict:
        """Parse an uploaded Excel file and create financial statements."""
        wb = load_workbook(io.BytesIO(file_content), data_only=True)

        stmt = select(StatementItemDict)
        result = await self.db.execute(stmt)
        all_items = result.scalars().all()
        item_by_code = {item.code: item for item in all_items}

        sheet_type_map = {
            "Баланс": "balance_sheet",
            "Отчёт о прибылях и убытках": "income_statement",
            "Отчёт о движении ДС": "cash_flow",
        }

        results = {"statements_created": 0, "items_created": 0, "errors": []}

        for sheet_name in wb.sheetnames:
            statement_type = sheet_type_map.get(sheet_name)
            if not statement_type:
                continue

            ws = wb[sheet_name]
            periods = self._detect_periods(ws)
            if not periods:
                results["errors"].append(f"No periods detected in sheet '{sheet_name}'")
                continue

            for period_col, period_end in periods.items():
                period_start = datetime.date(period_end.year, 1, 1)
                period_type = "annual" if period_end.month == 12 else "quarterly"

                fs = FinancialStatement(
                    company_id=company_id,
                    standard=standard,
                    statement_type=statement_type,
                    period_type=period_type,
                    period_start=period_start,
                    period_end=period_end,
                    currency="RUB",
                    unit_multiplier=1000,
                    source="manual_upload",
                )
                self.db.add(fs)
                await self.db.flush()
                results["statements_created"] += 1

                for row in ws.iter_rows(min_row=2, values_only=False):
                    code_cell = row[0]
                    value_cell = row[period_col - 1] if period_col - 1 < len(row) else None

                    if not code_cell or not code_cell.value:
                        continue

                    code = str(code_cell.value).strip()
                    dict_item = item_by_code.get(code)
                    if not dict_item:
                        continue

                    if value_cell and value_cell.value is not None:
                        try:
                            value = Decimal(str(value_cell.value))
                        except (InvalidOperation, ValueError):
                            continue

                        item = StatementItem(
                            statement_id=fs.id,
                            item_dict_id=dict_item.id,
                            value=value,
                        )
                        self.db.add(item)
                        results["items_created"] += 1

                await self.db.flush()

        wb.close()
        return results

    def _detect_periods(self, ws) -> dict[int, datetime.date]:
        """Detect period dates from header row."""
        import re

        periods = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            if col_idx < 3 or not cell.value:
                continue

            value = str(cell.value)
            date_match = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", value)
            if date_match:
                try:
                    d = datetime.date(
                        int(date_match.group(3)),
                        int(date_match.group(2)),
                        int(date_match.group(1)),
                    )
                    periods[col_idx] = d
                except ValueError:
                    pass

        return periods
